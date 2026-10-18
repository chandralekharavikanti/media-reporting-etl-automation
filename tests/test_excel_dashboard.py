"""Smoke tests for the generated executive workbook."""
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from src.build_excel_dashboard import build

ROOT = Path(__file__).resolve().parents[1]

class ExcelDashboardTests(unittest.TestCase):
    def test_chart_data_and_quality_count(self):
        with tempfile.TemporaryDirectory() as directory:
            path = build(report_dir=ROOT / "reports", output=Path(directory) / "dashboard.xlsx")
            wb = load_workbook(path)
            self.assertEqual(len(wb["Executive Dashboard"]._charts), 2)
            self.assertEqual(len(wb["Executive Dashboard"]._charts[0].series), 1)
            self.assertEqual(len(wb["Executive Dashboard"]._charts[1].series), 1)
            self.assertEqual(wb["Executive Dashboard"]["J6"].value, "=COUNTA('Data Quality'!A2:A1048576)")
            self.assertEqual(wb["Chart Data"].max_row, 13)
            self.assertIsInstance(wb["Chart Data"]["B2"].value, int)
            self.assertGreater(wb["Chart Data"]["B2"].value, 0)
            self.assertEqual(wb["Data Quality"]["E3"].value, "VALIDATION PASSED")
            for chart in wb["Executive Dashboard"]._charts:
                self.assertIn("Chart Data", str(chart.series[0].val.numRef.f))

if __name__ == "__main__":
    unittest.main()
