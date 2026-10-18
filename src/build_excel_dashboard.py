"""Generate a styled, reproducible Excel executive dashboard from ETL report CSVs.

Run: python -m src.build_excel_dashboard
Requires: openpyxl>=3.1
"""
import csv
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
NAVY = "14243B"
BLUE = "2165B5"
PALE = "EAF2FC"
TEAL = "087E8B"

def read_csv(path):
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))

def make_table(wb, title, headers, rows, money_col=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    for i, head in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(18, min(26, len(head) + 6))
    if rows:
        tab = Table(displayName=title.replace(" ", "") + "Table", ref=f"A1:{get_column_letter(len(headers))}{len(rows)+1}")
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tab)
    if money_col:
        for cells in ws.iter_cols(min_col=money_col, max_col=money_col, min_row=2):
            for cell in cells:
                cell.number_format = '"$"#,##0.00'
    return ws

def build(report_dir=None, output=None):
    report_dir = Path(report_dir or ROOT / "reports")
    output = Path(output or report_dir / "Media_Reporting_Executive_Dashboard.xlsx")
    monthly = read_csv(report_dir / "monthly_summary.csv")
    daily = read_csv(report_dir / "daily_traffic.csv")
    circulation = read_csv(report_dir / "monthly_circulation.csv")
    audit = read_csv(report_dir / "quality_audit.csv")
    wb = Workbook()
    ws = wb.active
    ws.title = "Executive Dashboard"
    ws.sheet_view.showGridLines = False

    media_rows = [[r["month"], r["publication"], int(r["sessions"]), int(r["pageviews"]), float(r["booked_revenue"])] for r in monthly]
    circ_rows = [[r["month"], r["publication"], int(r["circulation"])] for r in circulation]
    make_table(wb, "Monthly Media", ["Month", "Publication", "Sessions", "Pageviews", "Booked Revenue"], media_rows, 5)
    make_table(wb, "Daily Traffic", ["Date", "Sessions", "Pageviews"], [[r["date"], int(r["sessions"]), int(r["pageviews"])] for r in daily])
    make_table(wb, "Print Circulation", ["Month", "Publication", "Print Copies"], circ_rows)
    quality_sheet = make_table(wb, "Data Quality", ["Source", "Row or ID", "Finding"], [[r["source"], r["row_or_id"], r["issue"]] for r in audit])
    if not audit:
        quality_sheet["E3"] = "VALIDATION PASSED"
        quality_sheet["F3"] = "0 issues detected"
        quality_sheet["G3"] = "All source rows passed the configured checks."
        for cell in quality_sheet[3][4:7]:
            cell.fill = PatternFill("solid", fgColor="E8F5E9")
            cell.font = Font(color="176B39", bold=True)
        quality_sheet.column_dimensions["G"].width = 48

    ws.merge_cells("A1:M2")
    ws["A1"] = "MEDIA REPORTING  |  EXECUTIVE DASHBOARD"
    ws["A1"].font = Font(name="Aptos Display", size=19, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 19
    ws.merge_cells("A3:M3")
    ws["A3"] = "2025 synthetic publisher portfolio  •  automated Python + SQL ETL  •  Excel reporting"
    ws["A3"].font = Font(color=NAVY, italic=True, size=11)
    ws["A3"].fill = PatternFill("solid", fgColor=PALE)
    ws.row_dimensions[3].height = 27

    cards = [
        ("A5", "TOTAL SESSIONS", "A6", "=SUM('Monthly Media'!C2:C1048576)", "#,##0"),
        ("D5", "TOTAL PAGEVIEWS", "D6", "=SUM('Monthly Media'!D2:D1048576)", "#,##0"),
        ("G5", "BOOKED AD REVENUE", "G6", "=SUM('Monthly Media'!E2:E1048576)", '"$"#,##0'),
        ("J5", "QUALITY FINDINGS", "J6", "=COUNTA('Data Quality'!A2:A1048576)", "#,##0"),
    ]
    for label_cell, label, value_cell, formula, fmt in cards:
        ws[label_cell] = label
        ws[label_cell].fill = PatternFill("solid", fgColor=BLUE)
        ws[label_cell].font = Font(color="FFFFFF", bold=True, size=10)
        ws[value_cell] = formula
        ws[value_cell].font = Font(color=NAVY, bold=True, size=19)
        ws[value_cell].number_format = fmt
    ws.row_dimensions[5].height = 24
    ws.row_dimensions[6].height = 37

    ws["A9"] = "MONTHLY PERFORMANCE"
    ws["A9"].font = Font(bold=True, color=NAVY, size=13)
    for col, title in enumerate(["Month", "Sessions", "Pageviews", "Booked Revenue", "Print Copies"], 1):
        cell = ws.cell(10, col, title)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color="FFFFFF", bold=True)
    months = sorted({r[0] for r in media_rows} | {r[0] for r in circ_rows})
    for row_num, month in enumerate(months, 11):
        ws.cell(row_num, 1, month)
        for col, source_col in [(2,"C"), (3,"D"), (4,"E")]:
            ws.cell(row_num, col, f'=SUMIF(\'Monthly Media\'!$A$2:$A$1000,A{row_num},\'Monthly Media\'!${source_col}$2:${source_col}$1000)')
        ws.cell(row_num, 5, f'=SUMIF(\'Print Circulation\'!$A$2:$A$1000,A{row_num},\'Print Circulation\'!$C$2:$C$1000)')
        ws.cell(row_num, 4).number_format = '"$"#,##0'
        if row_num % 2:
            for cell in ws[row_num][:5]:
                cell.fill = PatternFill("solid", fgColor=PALE)

    # Chart data uses numeric values rather than formula cells: Excel for Mac and
    # non-Excel viewers can render it without waiting for formula recalculation.
    aggregates = defaultdict(lambda: [0, 0, 0, 0])
    for month, publication, sessions, pageviews, booked in media_rows:
        aggregates[month][0] += sessions
        aggregates[month][1] += pageviews
        aggregates[month][2] += booked
    for month, publication, copies in circ_rows:
        aggregates[month][3] += copies
    chart_data = make_table(
        wb, "Chart Data",
        ["Month", "Sessions", "Pageviews", "Booked Revenue", "Print Copies"],
        [[month] + aggregates[month] for month in months],
        4,
    )
    chart_data.sheet_properties.tabColor = TEAL
    end = len(months) + 1
    # Compact, readable charts: real numeric source data, one series each,
    # visible axis tick labels, consistent colors, and no misleading legends.
    categories = Reference(chart_data, min_col=1, min_row=2, max_row=end)
    traffic_chart = LineChart()
    traffic_chart.title = "Monthly Digital Sessions"
    traffic_chart.style = 13
    traffic_chart.add_data(Reference(chart_data, min_col=2, min_row=1, max_row=end), titles_from_data=True)
    traffic_chart.set_categories(categories)
    traffic_chart.legend = None
    traffic_chart.width, traffic_chart.height = 18.5, 9
    traffic_chart.y_axis.title = None
    traffic_chart.y_axis.numFmt = '#,##0'
    traffic_chart.y_axis.scaling.min = 0
    traffic_chart.y_axis.majorGridlines = None
    traffic_chart.y_axis.tickLblPos = "nextTo"
    traffic_chart.x_axis.tickLblPos = "nextTo"
    traffic_chart.x_axis.tickMarkSkip = 1
    traffic_chart.x_axis.lblOffset = 100
    traffic_chart.x_axis.crosses = "min"
    traffic_chart.series[0].graphicalProperties.line.solidFill = BLUE
    traffic_chart.series[0].graphicalProperties.line.width = 32000
    traffic_chart.series[0].smooth = False
    traffic_chart.display_blanks = "gap"
    ws.add_chart(traffic_chart, "G9")

    revenue_chart = BarChart()
    revenue_chart.type = "col"
    revenue_chart.grouping = "clustered"
    revenue_chart.overlap = 0
    revenue_chart.title = "Monthly Booked Advertising Revenue"
    revenue_chart.style = 10
    revenue_chart.add_data(Reference(chart_data, min_col=4, min_row=1, max_row=end), titles_from_data=True)
    revenue_chart.set_categories(categories)
    revenue_chart.legend = None
    revenue_chart.width, revenue_chart.height = 18.5, 9
    revenue_chart.y_axis.title = None
    revenue_chart.y_axis.numFmt = '"$"#,##0'
    revenue_chart.y_axis.scaling.min = 0
    revenue_chart.y_axis.majorGridlines = None
    revenue_chart.y_axis.tickLblPos = "nextTo"
    revenue_chart.x_axis.tickLblPos = "nextTo"
    revenue_chart.x_axis.crosses = "min"
    revenue_chart.series[0].graphicalProperties.solidFill = TEAL
    revenue_chart.series[0].graphicalProperties.line.noFill = True
    revenue_chart.gapWidth = 65
    ws.add_chart(revenue_chart, "G28")

    ws["A26"] = "PUBLICATION COMPARISON"
    ws["A26"].font = Font(bold=True, color=NAVY, size=13)
    for col, title in enumerate(["Publication", "Sessions", "Booked Revenue", "Print Copies"], 1):
        cell = ws.cell(27, col, title)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color="FFFFFF", bold=True)
    publications = sorted({r[1] for r in media_rows})
    for row_num, publication in enumerate(publications, 28):
        ws.cell(row_num, 1, publication)
        ws.cell(row_num, 2, f'=SUMIF(\'Monthly Media\'!$B$2:$B$1000,A{row_num},\'Monthly Media\'!$C$2:$C$1000)')
        ws.cell(row_num, 3, f'=SUMIF(\'Monthly Media\'!$B$2:$B$1000,A{row_num},\'Monthly Media\'!$E$2:$E$1000)')
        ws.cell(row_num, 4, f'=SUMIF(\'Print Circulation\'!$B$2:$B$1000,A{row_num},\'Print Circulation\'!$C$2:$C$1000)')
        ws.cell(row_num, 3).number_format = '"$"#,##0'
    ws["A35"] = "Notes: synthetic portfolio. Print copies are not digital impressions. Booked revenue is not recognized revenue."
    ws["A35"].font = Font(italic=True, color=TEAL, size=10)
    for col in "ABCDEFGHIJKLM":
        ws.column_dimensions[col].width = 17 if col != "A" else 23
    ws.freeze_panes = "A11"
    ws.sheet_view.zoomScale = 85
    ws.print_options.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.print_area = "A1:M47"
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return output

if __name__ == "__main__":
    print("Excel dashboard:", build())
